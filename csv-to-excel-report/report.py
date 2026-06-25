"""
CSV → Excel Report Generator
============================

Reads a raw CSV file, cleans it, and produces a polished multi-sheet Excel
report:

  • "Data"    — the cleaned dataset with a styled header and auto-sized columns
  • "Summary" — totals of a numeric column grouped by a category column,
                plus an auto-generated bar chart

Usage
-----
    python report.py sales.csv
    python report.py sales.csv -o report.xlsx --group-by Region --value Revenue

If --group-by / --value are omitted, the script auto-detects a sensible
text column and numeric column.
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Ensure emoji / Unicode print correctly on Windows consoles (cp1251 etc.)
try:
    sys.stdout.reconfigure(encoding="utf-8")
except (AttributeError, ValueError):
    pass

HEADER_FILL = PatternFill("solid", fgColor="2F5496")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        sys.exit(f"❌ File not found: {path}")
    try:
        return pd.read_csv(path)
    except Exception as exc:  # noqa: BLE001
        sys.exit(f"❌ Could not read CSV: {exc}")


def clean(df: pd.DataFrame) -> pd.DataFrame:
    """Trim column names, drop empty rows/columns, strip string cells."""
    df = df.copy()
    df.columns = [str(c).strip() for c in df.columns]
    df = df.dropna(axis=0, how="all").dropna(axis=1, how="all")
    for col in df.select_dtypes(include="object"):
        df[col] = df[col].astype(str).str.strip()
    return df.reset_index(drop=True)


def pick_columns(
    df: pd.DataFrame, group_by: str | None, value: str | None
) -> tuple[str, str]:
    """Resolve / auto-detect the category and numeric columns."""
    numeric_cols = df.select_dtypes(include="number").columns.tolist()
    text_cols = [c for c in df.columns if c not in numeric_cols]

    if value is None:
        if not numeric_cols:
            sys.exit("❌ No numeric column found to summarize. Use --value.")
        value = numeric_cols[0]
    elif value not in df.columns:
        sys.exit(f"❌ --value column '{value}' not found. Columns: {list(df.columns)}")

    if group_by is None:
        group_by = text_cols[0] if text_cols else df.columns[0]
    elif group_by not in df.columns:
        sys.exit(f"❌ --group-by column '{group_by}' not found. Columns: {list(df.columns)}")

    return group_by, value


def style_header(ws: Worksheet) -> None:
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def autosize(ws: Worksheet) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(width + 2, 50)


def build_report(
    df: pd.DataFrame, out: Path, group_by: str, value: str
) -> pd.DataFrame:
    summary = (
        df.groupby(group_by, dropna=False)[value]
        .sum()
        .sort_values(ascending=False)
        .reset_index()
    )

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Data", index=False)
        summary.to_excel(writer, sheet_name="Summary", index=False)

        wb = writer.book
        data_ws = writer.sheets["Data"]
        sum_ws = writer.sheets["Summary"]

        for ws in (data_ws, sum_ws):
            style_header(ws)
            ws.freeze_panes = "A2"
            autosize(ws)

        # Bar chart of the summary
        chart = BarChart()
        chart.title = f"{value} by {group_by}"
        chart.y_axis.title = value
        chart.x_axis.title = group_by
        n = len(summary)
        data_ref = Reference(sum_ws, min_col=2, min_row=1, max_row=n + 1)
        cats_ref = Reference(sum_ws, min_col=1, min_row=2, max_row=n + 1)
        chart.add_data(data_ref, titles_from_data=True)
        chart.set_categories(cats_ref)
        chart.height = 8
        chart.width = 16
        anchor_col = get_column_letter(summary.shape[1] + 2)
        sum_ws.add_chart(chart, f"{anchor_col}2")

    return summary


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Generate a formatted Excel report from a CSV file."
    )
    parser.add_argument("csv", type=Path, help="Path to the input CSV file")
    parser.add_argument(
        "-o", "--output", type=Path, help="Output .xlsx path (default: <csv>_report.xlsx)"
    )
    parser.add_argument("--group-by", help="Category column to group by (auto-detected if omitted)")
    parser.add_argument("--value", help="Numeric column to summarize (auto-detected if omitted)")
    args = parser.parse_args()

    df = clean(load_csv(args.csv))
    if df.empty:
        sys.exit("❌ The CSV has no usable data after cleaning.")

    group_by, value = pick_columns(df, args.group_by, args.value)
    out = args.output or args.csv.with_name(f"{args.csv.stem}_report.xlsx")

    summary = build_report(df, out, group_by, value)

    print(f"✅ Report written to: {out}")
    print(f"   Rows processed : {len(df)}")
    print(f"   Grouped by     : {group_by}")
    print(f"   Summed         : {value}")
    print(f"   Categories     : {len(summary)}")


if __name__ == "__main__":
    main()
